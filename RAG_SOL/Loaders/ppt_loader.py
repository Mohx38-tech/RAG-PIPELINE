from pathlib import Path
from pptx import Presentation


def load_pptx(file_path: str | Path) -> str:
    file_path = Path(file_path)
    presentation = Presentation(str(file_path))

    slides_text = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text_parts.append(shape.text)

        if slide_text_parts:
            slide_text = " ".join(slide_text_parts)
            slides_text.append(f"\n--- Slide {slide_number} ---\n{slide_text}")

    return "\n".join(slides_text)